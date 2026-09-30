"""Write organizer sessions to a Word document.

Real images are embedded. A file that is not an image (the sample set includes
one text placeholder) is written as its filename, and every screenshot gets a
caption with the filename and the timestamp parsed from that name.
"""

from __future__ import annotations

from pathlib import Path

from model import Session, Shot


# Signatures python-docx knows how to embed. Checked before add_picture so a
# text placeholder does not have to fail inside the library first.
_IMAGE_SIGNATURES = (
    b"\x89PNG\r\n\x1a\n",
    b"\xff\xd8\xff",
    b"GIF87a",
    b"GIF89a",
    b"BM",
    b"II*\x00",
    b"MM\x00*",
)


def is_embeddable_image(path: Path) -> bool:
    try:
        header = path.read_bytes()[:8]
    except OSError:
        return False
    return any(header.startswith(signature) for signature in _IMAGE_SIGNATURES)


def export_document(sessions: list[Session], destination: Path) -> None:
    """Create one document: a heading per session and a subheading per group."""

    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is not installed. From the repository root, run: "
            "pip install -r requirements.txt"
        ) from exc

    document = Document()
    picture_width = Inches(5.5)

    for session in sessions:
        document.add_heading(session.name, level=1)
        if not session.groups:
            document.add_paragraph("No buckets in this session.")
            continue
        for group in session.groups:
            document.add_heading(group.label, level=2)
            if not group.shots:
                document.add_paragraph("No screenshots in this bucket.")
                continue
            for shot in group.shots:
                _add_shot(document, shot, picture_width)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))


def _add_shot(document: object, shot: Shot, picture_width: object) -> None:
    if is_embeddable_image(shot.path):
        try:
            document.add_picture(str(shot.path), width=picture_width)  # type: ignore[attr-defined]
        except Exception:
            # The header looked like an image, but the file could not be embedded.
            document.add_paragraph(shot.filename)  # type: ignore[attr-defined]
    else:
        document.add_paragraph(shot.filename)  # type: ignore[attr-defined]

    # Caption style is part of python-docx's default template.
    document.add_paragraph(shot.caption(), style="Caption")  # type: ignore[attr-defined]
